import json
from typing import Any, Dict, List

from pydantic import BaseModel, Field

from app.sandbox.limits import check_file_size, resolve_sandbox_path
from app.tools.registry import tool_registry


def _ok(message: str, data: dict = None, files_modified: list = None) -> str:
    return json.dumps(
        {
            "status": "success",
            "message": message,
            "data": data or {},
            "files_modified": files_modified or [],
        }
    )


def _err(message: str) -> str:
    return json.dumps({"status": "error", "message": message})


class CreateXlsxFromRowsArgs(BaseModel):
    path: str = Field(..., description="Relative path inside sandbox ending with .xlsx")
    rows: List[Dict[str, Any]] = Field(
        ..., description="List of row objects to write into the sheet"
    )
    columns: List[str] = Field(
        default_factory=lambda: ["Rank", "Headline", "Story", "Source", "Published At"],
        description="Ordered list of columns to write as headers",
    )
    sheet_name: str = Field("LatestNews", description="Excel sheet name")


@tool_registry.register(
    "create_xlsx_from_rows",
    "Create a valid .xlsx spreadsheet from structured rows with explicit column ordering.",
    CreateXlsxFromRowsArgs,
)
def create_xlsx_from_rows(args: CreateXlsxFromRowsArgs) -> str:
    if not args.path.lower().endswith(".xlsx"):
        return _err("Path must end with .xlsx")

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except Exception:
        return _err("openpyxl is not installed. Add it to requirements and install dependencies.")

    try:
        path = resolve_sandbox_path(args.path)
        path.parent.mkdir(parents=True, exist_ok=True)

        wb = Workbook()
        ws = wb.active
        ws.title = (args.sheet_name or "LatestNews")[:31]

        columns = args.columns or ["Rank", "Headline", "Story", "Source", "Published At"]

        # Header
        for col_idx, col_name in enumerate(columns, start=1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.font = Font(bold=True)

        # Rows
        for row_idx, row in enumerate(args.rows, start=2):
            row_data = dict(row)
            if "Rank" in columns and (row_data.get("Rank") in (None, "")):
                row_data["Rank"] = row_idx - 1

            for col_idx, col_name in enumerate(columns, start=1):
                ws.cell(row=row_idx, column=col_idx, value=row_data.get(col_name, ""))

        # Basic column sizing for readability.
        for col_idx, col_name in enumerate(columns, start=1):
            max_len = len(str(col_name))
            for row_idx in range(2, len(args.rows) + 2):
                value = ws.cell(row=row_idx, column=col_idx).value
                if value is None:
                    continue
                max_len = max(max_len, len(str(value)))
            ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = min(
                max_len + 2, 80
            )

        wb.save(str(path))
        check_file_size(path)

        return _ok(
            f"Created xlsx file at {args.path}",
            data={"rows_written": len(args.rows), "columns": columns},
            files_modified=[args.path],
        )
    except Exception as exc:
        return _err(str(exc))
