import json
import csv
import re
import xml.etree.ElementTree as ET
from datetime import datetime
from difflib import SequenceMatcher
from html import unescape
from io import StringIO
from pathlib import Path
from typing import Any, Dict, List, Optional
from urllib.parse import urlparse
from xml.sax.saxutils import escape
import zipfile

import requests
from pydantic import BaseModel, Field

from app.config import settings
from app.tools.registry import tool_registry


CATEGORY_FEEDS: Dict[str, List[str]] = {
    "sports": [
        "https://feeds.reuters.com/reuters/sportsNews",
        "http://rss.cnn.com/rss/edition_sport.rss",
        "https://feeds.bbci.co.uk/sport/rss.xml",
    ],
    "entertainment": [
        "https://feeds.reuters.com/reuters/entertainment",
        "http://rss.cnn.com/rss/edition_entertainment.rss",
        "https://feeds.bbci.co.uk/news/entertainment_and_arts/rss.xml",
    ],
    "business": [
        "https://feeds.reuters.com/reuters/businessNews",
        "http://rss.cnn.com/rss/money_latest.rss",
        "https://feeds.bbci.co.uk/news/business/rss.xml",
    ],
    "technology": [
        "https://feeds.reuters.com/reuters/technologyNews",
        "http://rss.cnn.com/rss/edition_technology.rss",
        "https://feeds.bbci.co.uk/news/technology/rss.xml",
    ],
    "health": [
        "https://feeds.reuters.com/reuters/healthNews",
        "http://rss.cnn.com/rss/edition_health.rss",
        "https://feeds.bbci.co.uk/news/health/rss.xml",
    ],
}


BOILERPLATE_MARKERS = (
    "close icon",
    "download the cnn app",
    "sign in to your cnn account",
    "my account settings",
    "newsletters",
    "topics you follow",
    "terms of use",
    "privacy policy",
    "d=\"m24",
    "style=\"fill:",
)


def _ok(message: str, data: dict = None, files_modified: list = None) -> str:
    return json.dumps(
        {
            "status": "success",
            "message": message,
            "data": data or {},
            "files_modified": files_modified or [],
        }
    )


def _err(message: str) -> str:
    return json.dumps({"status": "error", "message": message})


def _local_name(tag: str) -> str:
    if "}" in tag:
        return tag.split("}", 1)[1]
    return tag


def _clean_text(value: str) -> str:
    value = unescape(value or "")
    value = re.sub(r"<[^>]+>", " ", value)
    value = re.sub(r"\s+", " ", value)
    return value.strip()


def _choose_better_story(current_story: str, candidate_story: str) -> str:
    current = (current_story or "").strip()
    candidate = (candidate_story or "").strip()
    if not candidate:
        return current
    if not current:
        return candidate
    if len(candidate.split()) >= len(current.split()) + 12 and not _looks_like_boilerplate(candidate):
        return candidate
    return current


def _normalize_topic(topic: str) -> str:
    if not topic:
        return settings.NEWS_CANONICAL_TOPIC.strip()
    low = topic.lower()
    if "iran" in low and any(token in low for token in ("israel", "us", "u.s", "america")):
        return settings.NEWS_CANONICAL_TOPIC.strip() or topic.strip()
    return topic.strip()


def _detect_category(topic: str) -> Optional[str]:
    low = (topic or "").casefold()
    category_keywords = {
        "sports": ("sport", "sports", "football", "soccer", "cricket", "nba", "tennis"),
        "entertainment": ("entertainment", "movie", "movies", "music", "celebrity", "showbiz"),
        "business": ("business", "finance", "stock", "market", "economy"),
        "technology": ("technology", "tech", "ai", "software", "hardware", "startup"),
        "health": ("health", "medical", "medicine", "hospital", "disease"),
    }
    for category, keys in category_keywords.items():
        if any(key in low for key in keys):
            return category
    return None


def _select_feeds(topic: str, feed_urls: Optional[List[str]]) -> List[str]:
    if feed_urls:
        return feed_urls

    category = _detect_category(topic)
    if category and CATEGORY_FEEDS.get(category):
        return list(CATEGORY_FEEDS[category])

    return settings.NEWS_RSS_FEEDS


def _story_is_thin(headline: str, story: str) -> bool:
    h = (headline or "").strip()
    s = (story or "").strip()
    if not s:
        return True
    if len(s.split()) < 12:
        return True
    if h:
        similarity = SequenceMatcher(None, h.casefold(), s.casefold()).ratio()
        if similarity >= 0.72:
            return True
    return False


def _looks_like_boilerplate(text: str) -> bool:
    t = (text or "").strip().casefold()
    if not t:
        return True
    if any(marker in t for marker in BOILERPLATE_MARKERS):
        return True

    words = t.split()
    punctuation = sum(1 for ch in t if ch in ".!?")
    if len(words) > 180 and punctuation < 3:
        return True

    repeated_menu_tokens = sum(1 for tok in ("markets", "tech", "media", "videos", "watch", "listen") if tok in t)
    if repeated_menu_tokens >= 4 and len(words) > 80:
        return True

    return False


def _headline_is_noise(headline: str) -> bool:
    h = (headline or "").strip().casefold()
    if not h:
        return True
    noise_phrases = (
        "sign up to our newsletter",
        "download the cnn app",
        "something isn't loading properly",
        "ad feedback",
    )
    return any(p in h for p in noise_phrases)


def _extract_story_from_html(html_text: str) -> str:
    if not html_text:
        return ""

    cleaned = re.sub(r"<script[\s\S]*?</script>", " ", html_text, flags=re.IGNORECASE)
    cleaned = re.sub(r"<style[\s\S]*?</style>", " ", cleaned, flags=re.IGNORECASE)

    json_ld_blobs = re.findall(
        r'<script[^>]+type=["\']application/ld\+json["\'][^>]*>([\s\S]*?)</script>',
        html_text,
        flags=re.IGNORECASE,
    )
    for blob in json_ld_blobs:
        raw = (blob or "").strip()
        if not raw:
            continue
        try:
            payload = json.loads(raw)
        except Exception:
            continue

        stack = payload if isinstance(payload, list) else [payload]
        while stack:
            item = stack.pop()
            if isinstance(item, list):
                stack.extend(item)
                continue
            if not isinstance(item, dict):
                continue

            for key in ("articleBody", "text", "description"):
                candidate = _clean_text(str(item.get(key, "")))
                if len(candidate.split()) >= 80 and not _looks_like_boilerplate(candidate):
                    words = candidate.split()
                    if len(words) > 500:
                        candidate = " ".join(words[:500]) + " ..."
                    return candidate

            for key in ("@graph", "hasPart", "mainEntity"):
                nested = item.get(key)
                if nested is not None:
                    stack.append(nested)

    article_blocks = re.findall(r"<article[^>]*>([\s\S]*?)</article>", cleaned, flags=re.IGNORECASE)
    for block in article_blocks:
        pieces = []
        for p in re.findall(r"<p[^>]*>([\s\S]*?)</p>", block, flags=re.IGNORECASE):
            text = _clean_text(p)
            if len(text) >= 40 and not _looks_like_boilerplate(text):
                pieces.append(text)
        if pieces:
            article_text = " ".join(pieces)
            if len(article_text.split()) >= 80:
                words = article_text.split()
                if len(words) > 500:
                    article_text = " ".join(words[:500]) + " ..."
                return article_text

    paragraphs: List[str] = []
    for p in re.findall(r"<p[^>]*>([\s\S]*?)</p>", cleaned, flags=re.IGNORECASE):
        text = _clean_text(p)
        if len(text) >= 40 and not _looks_like_boilerplate(text):
            paragraphs.append(text)
        if len(paragraphs) >= 10:
            break

    if paragraphs:
        merged = " ".join(paragraphs)
        words = merged.split()
        if len(words) > 500:
            merged = " ".join(words[:500]) + " ..."
        return merged

    meta_match = re.search(
        r'<meta[^>]+(?:name|property)=["\'](?:description|og:description)["\'][^>]*content=["\']([^"\']+)["\']',
        cleaned,
        flags=re.IGNORECASE,
    )
    if meta_match:
        meta_text = _clean_text(meta_match.group(1))
        if not _looks_like_boilerplate(meta_text):
            return meta_text

    return ""


def _try_fetch_article_story(url: str) -> str:
    parsed = urlparse(url)
    if parsed.scheme not in ("http", "https"):
        return ""

    response = requests.get(
        url,
        timeout=max(3, min(8, settings.NEWS_REQUEST_TIMEOUT_SECONDS)),
        headers={"User-Agent": "MCPLinkV2-NewsFetcher/1.0 (+https://local-agent)"},
    )
    response.raise_for_status()

    content_type = (response.headers or {}).get("Content-Type", "")
    if "text/html" not in content_type and "application/xhtml+xml" not in content_type:
        return ""

    return _extract_story_from_html(response.text)


def _parse_timestamp(value: str) -> str:
    if not value:
        return ""
    raw = value.strip()
    for fmt in (
        "%a, %d %b %Y %H:%M:%S %z",
        "%a, %d %b %Y %H:%M:%S %Z",
        "%Y-%m-%dT%H:%M:%S%z",
        "%Y-%m-%dT%H:%M:%SZ",
        "%Y-%m-%dT%H:%M:%S.%fZ",
    ):
        try:
            dt = datetime.strptime(raw, fmt)
            return dt.isoformat()
        except ValueError:
            continue
    return raw


def _extract_items(xml_bytes: bytes, feed_url: str) -> List[Dict[str, str]]:
    try:
        root = ET.fromstring(xml_bytes)
    except ET.ParseError:
        return []

    host = urlparse(feed_url).hostname or ""
    source_name = host
    channel_title = root.find("./channel/title")
    if channel_title is not None and channel_title.text:
        source_name = channel_title.text.strip()

    rows: List[Dict[str, str]] = []

    # RSS 2.0
    for item in root.findall(".//item"):
        title = _clean_text(item.findtext("title", ""))
        if _headline_is_noise(title):
            continue
        story = _clean_text(item.findtext("description", ""))

        for child in list(item):
            local = _local_name(child.tag)
            if local in ("content", "encoded", "summary", "content:encoded"):
                story = _choose_better_story(story, _clean_text(child.text or ""))

        link = _clean_text(item.findtext("link", ""))
        published = _parse_timestamp(item.findtext("pubDate", ""))
        if title:
            if _headline_is_noise(title):
                continue
            rows.append(
                {
                    "Headline": title,
                    "Story": story,
                    "Source": source_name,
                    "Published At": published,
                    "URL": link,
                    "Story Source": "rss",
                }
            )

    # Atom
    for entry in root.iter():
        if _local_name(entry.tag) != "entry":
            continue

        title = ""
        story = ""
        published = ""
        link = ""

        for child in list(entry):
            name = _local_name(child.tag)
            if name == "title":
                title = _clean_text(child.text or "")
            elif name in ("summary", "content"):
                story = _choose_better_story(story, _clean_text(child.text or ""))
            elif name in ("updated", "published"):
                published = _parse_timestamp(child.text or "")
            elif name == "link":
                href = (child.attrib or {}).get("href", "")
                if href:
                    link = _clean_text(href)

        if title:
            if _headline_is_noise(title):
                continue
            rows.append(
                {
                    "Headline": title,
                    "Story": story,
                    "Source": source_name,
                    "Published At": published,
                    "URL": link,
                    "Story Source": "rss",
                }
            )

    return rows


def _score_relevance(text: str, keywords: List[str]) -> int:
    if not text:
        return 0
    low = text.lower()
    return sum(1 for kw in keywords if kw in low)


class FetchRssNewsArgs(BaseModel):
    topic: str = Field(..., description="Requested news topic or conflict phrase")
    limit: int = Field(10, ge=1, le=50, description="Maximum number of rows to return")
    feed_urls: Optional[List[str]] = Field(
        None,
        description="Optional RSS/Atom feed URLs. Defaults to configured sources.",
    )
    sort_mode: str = Field(
        "headline_alphabetical",
        description="Sorting mode. Use 'headline_alphabetical' for A-Z headline ordering.",
    )
    story_mode: str = Field(
        "expanded",
        description="Story handling mode: 'expanded' to enrich thin stories, 'rss_only' to keep raw feed stories.",
    )


def _fetch_news_rows(
    topic: str,
    requested_limit: int,
    feed_urls: Optional[List[str]],
    sort_mode: str,
    story_mode: str = "expanded",
) -> Dict[str, Any]:
    limit = min(requested_limit, settings.NEWS_MAX_ITEMS_PER_REQUEST)
    feeds = _select_feeds(topic, feed_urls)
    if not feeds:
        raise ValueError("No RSS feeds configured.")

    canonical_topic = _normalize_topic(topic)
    topic_for_matching = canonical_topic.lower()
    keywords = [
        token
        for token in re.findall(r"[a-zA-Z]{3,}", topic_for_matching)
        if token not in {"with", "from", "that", "this", "ongoing", "conflict"}
    ]

    all_rows: List[Dict[str, str]] = []
    errors: List[str] = []

    headers = {
        "User-Agent": "MCPLinkV2-NewsFetcher/1.0 (+https://local-agent)",
        "Accept": "application/rss+xml, application/atom+xml, application/xml, text/xml",
    }

    for feed_url in feeds:
        try:
            response = requests.get(
                feed_url,
                timeout=settings.NEWS_REQUEST_TIMEOUT_SECONDS,
                headers=headers,
            )
            response.raise_for_status()
            all_rows.extend(_extract_items(response.content, feed_url))
        except Exception as exc:
            errors.append(f"{feed_url}: {exc}")

    if not all_rows:
        details = "; ".join(errors) if errors else "Unknown fetch/parsing error"
        raise ValueError(f"Unable to fetch any news rows. Details: {details}")

    deduped: List[Dict[str, str]] = []
    seen = set()
    for row in all_rows:
        key = row.get("Headline", "").strip().lower()
        if not key or key in seen:
            continue
        seen.add(key)
        deduped.append(row)

    ranked = [
        (
            _score_relevance(
                f"{row.get('Headline', '')} {row.get('Story', '')}",
                keywords,
            ),
            row,
        )
        for row in deduped
    ]
    relevant = [row for score, row in ranked if score > 0]
    category = _detect_category(topic)
    if keywords and not category:
        pool = relevant
    else:
        pool = relevant if relevant else deduped

    if sort_mode == "headline_alphabetical":
        pool.sort(key=lambda row: row.get("Headline", "").casefold())

    selected = pool[:limit]

    enrichment_errors = 0
    if story_mode != "rss_only":
        for row in selected:
            if _story_is_thin(row.get("Headline", ""), row.get("Story", "")):
                try:
                    enriched = _try_fetch_article_story(row.get("URL", ""))
                    if enriched and not _looks_like_boilerplate(enriched):
                        row["Story"] = enriched
                        row["Story Source"] = "article_html"
                except Exception:
                    enrichment_errors += 1

            if _story_is_thin(row.get("Headline", ""), row.get("Story", "")):
                url_hint = f" Full article: {row.get('URL')}" if row.get("URL") else ""
                row["Story"] = (
                    "Detailed story text was not fully available in the feed. "
                    "This item is included based on headline relevance and source metadata."
                    f"{url_hint}"
                )
                row["Story Source"] = "fallback"

    for idx, row in enumerate(selected, start=1):
        row["Rank"] = idx

    source_counts: Dict[str, int] = {}
    story_word_counts: List[int] = []
    for row in selected:
        src = row.get("Story Source", "unknown")
        source_counts[src] = source_counts.get(src, 0) + 1
        story_word_counts.append(len((row.get("Story", "") or "").split()))

    shortage = max(0, limit - len(selected))
    return {
        "requested_topic": topic,
        "normalized_topic": canonical_topic,
        "requested_limit": requested_limit,
        "returned_count": len(selected),
        "shortage": shortage,
        "rows": selected,
        "partial_source_errors": errors,
        "story_mode": story_mode,
        "story_enrichment_errors": enrichment_errors,
        "story_source_counts": source_counts,
        "story_depth": {
            "avg_words": (sum(story_word_counts) / len(story_word_counts)) if story_word_counts else 0,
            "min_words": min(story_word_counts) if story_word_counts else 0,
            "max_words": max(story_word_counts) if story_word_counts else 0,
        },
    }


@tool_registry.register(
    "fetch_rss_news",
    "Fetch latest news from trusted RSS feeds and return normalized rows suitable for spreadsheet export.",
    FetchRssNewsArgs,
)
def fetch_rss_news(args: FetchRssNewsArgs) -> str:
    try:
        data = _fetch_news_rows(
            args.topic,
            args.limit,
            args.feed_urls,
            args.sort_mode,
            args.story_mode,
        )
        message = f"Fetched {data['returned_count']} news rows"
        if data["shortage"]:
            message += f" (short by {data['shortage']} from requested {min(args.limit, settings.NEWS_MAX_ITEMS_PER_REQUEST)})."
        else:
            message += f" for requested {min(args.limit, settings.NEWS_MAX_ITEMS_PER_REQUEST)}."

        return _ok(
            message,
            data=data,
        )
    except Exception as exc:
        return _err(str(exc))


class CreateLatestNewsXlsxArgs(BaseModel):
    path: str = Field(..., description="Relative path for output .xlsx in sandbox")
    topic: str = Field(..., description="News topic phrase to fetch")
    limit: int = Field(10, ge=1, le=50, description="Maximum number of rows to write")
    feed_urls: Optional[List[str]] = Field(
        None,
        description="Optional RSS/Atom feed URLs. Defaults to configured sources.",
    )
    sort_mode: str = Field(
        "headline_alphabetical",
        description="Sorting mode. Use 'headline_alphabetical' for A-Z headline ordering.",
    )
    sheet_name: str = Field("LatestNews", description="Excel sheet name")
    columns: List[str] = Field(
        default_factory=lambda: ["Rank", "Headline", "Story", "Source", "Published At"],
        description="Ordered list of columns to write",
    )
    story_mode: str = Field(
        "expanded",
        description="Story handling mode: 'expanded' to enrich thin stories, 'rss_only' to keep raw feed stories.",
    )


@tool_registry.register(
    "create_latest_news_xlsx",
    "Fetch latest news from RSS feeds and write directly into a .xlsx file in one deterministic step.",
    CreateLatestNewsXlsxArgs,
)
def create_latest_news_xlsx(args: CreateLatestNewsXlsxArgs) -> str:
    if not args.path.lower().endswith(".xlsx"):
        return _err("Path must end with .xlsx")

    try:
        universal_args = CreateLatestNewsFileArgs(
            path=args.path,
            topic=args.topic,
            limit=args.limit,
            feed_urls=args.feed_urls,
            sort_mode=args.sort_mode,
            story_mode=args.story_mode,
            sheet_name=args.sheet_name,
            columns=args.columns,
        )
        return create_latest_news_file(universal_args)
    except Exception as exc:
        return _err(str(exc))


class CreateLatestNewsFileArgs(BaseModel):
    path: str = Field(..., description="Relative output path in sandbox (.xlsx, .txt, .docx, .md, .json, .csv)")
    topic: str = Field(..., description="News topic phrase to fetch (sports, entertainment, politics, etc.)")
    limit: int = Field(10, ge=1, le=50, description="Maximum number of news rows to write")
    feed_urls: Optional[List[str]] = Field(
        None,
        description="Optional RSS/Atom feed URLs. Defaults to configured sources.",
    )
    sort_mode: str = Field(
        "headline_alphabetical",
        description="Sorting mode. Use 'headline_alphabetical' for A-Z headline ordering.",
    )
    story_mode: str = Field(
        "expanded",
        description="Story handling mode: 'expanded' to enrich thin stories, 'rss_only' to keep raw feed stories.",
    )
    sheet_name: str = Field("LatestNews", description="Used only for xlsx output")
    columns: List[str] = Field(
        default_factory=lambda: ["Rank", "Headline", "Story", "Source", "Published At"],
        description="Column order for structured formats",
    )


def _write_xlsx(path: Path, rows: List[Dict[str, Any]], columns: List[str], sheet_name: str) -> None:
    def _col_letter(idx: int) -> str:
        result = ""
        while idx > 0:
            idx, rem = divmod(idx - 1, 26)
            result = chr(65 + rem) + result
        return result

    def _cell_xml(ref: str, value: Any) -> str:
        if value is None:
            return f'<c r="{ref}" t="inlineStr"><is><t></t></is></c>'
        if isinstance(value, (int, float)):
            return f'<c r="{ref}"><v>{value}</v></c>'
        text = escape(str(value))
        return f'<c r="{ref}" t="inlineStr"><is><t>{text}</t></is></c>'

    sheet_title = escape((sheet_name or "LatestNews")[:31])

    all_rows: List[List[Any]] = [columns]
    for i, row in enumerate(rows, start=1):
        row_data = dict(row)
        if "Rank" in columns and (row_data.get("Rank") in (None, "")):
            row_data["Rank"] = i
        all_rows.append([row_data.get(col, "") for col in columns])

    sheet_rows_xml: List[str] = []
    for r_idx, row_values in enumerate(all_rows, start=1):
        cells = []
        for c_idx, value in enumerate(row_values, start=1):
            ref = f"{_col_letter(c_idx)}{r_idx}"
            cells.append(_cell_xml(ref, value))
        sheet_rows_xml.append(f"<row r=\"{r_idx}\">{''.join(cells)}</row>")

    worksheet_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        '<sheetData>'
        f"{''.join(sheet_rows_xml)}"
        '</sheetData>'
        '</worksheet>'
    )

    workbook_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        '<sheets>'
        f'<sheet name="{sheet_title}" sheetId="1" r:id="rId1"/>'
        '</sheets>'
        '</workbook>'
    )

    workbook_rels_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" '
        'Target="worksheets/sheet1.xml"/>'
        '<Relationship Id="rId2" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" '
        'Target="styles.xml"/>'
        '</Relationships>'
    )

    styles_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        '<fonts count="1"><font><sz val="11"/><name val="Calibri"/></font></fonts>'
        '<fills count="1"><fill><patternFill patternType="none"/></fill></fills>'
        '<borders count="1"><border/></borders>'
        '<cellStyleXfs count="1"><xf/></cellStyleXfs>'
        '<cellXfs count="1"><xf xfId="0"/></cellXfs>'
        '</styleSheet>'
    )

    root_rels_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
        'Target="xl/workbook.xml"/>'
        '</Relationships>'
    )

    content_types_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/xl/workbook.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
        '<Override PartName="/xl/worksheets/sheet1.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
        '<Override PartName="/xl/styles.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>'
        '</Types>'
    )

    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", content_types_xml)
        zf.writestr("_rels/.rels", root_rels_xml)
        zf.writestr("xl/workbook.xml", workbook_xml)
        zf.writestr("xl/_rels/workbook.xml.rels", workbook_rels_xml)
        zf.writestr("xl/styles.xml", styles_xml)
        zf.writestr("xl/worksheets/sheet1.xml", worksheet_xml)


def _write_text(path: Path, rows: List[Dict[str, Any]], topic: str) -> None:
    lines = [f"Top latest news for: {topic}", ""]
    for row in rows:
        lines.append(f"{row.get('Rank', '')}. {row.get('Headline', '')}")
        lines.append(f"Story: {row.get('Story', '')}")
        lines.append(f"Source: {row.get('Source', '')}")
        lines.append(f"Published At: {row.get('Published At', '')}")
        if row.get("URL"):
            lines.append(f"URL: {row.get('URL')}")
        lines.append("")
    path.write_text("\n".join(lines).strip() + "\n", encoding="utf-8")


def _write_docx(path: Path, rows: List[Dict[str, Any]], topic: str) -> None:
    from docx import Document

    doc = Document()
    doc.add_heading(f"Top latest news for: {topic}", level=1)

    for row in rows:
        doc.add_heading(f"{row.get('Rank', '')}. {row.get('Headline', '')}", level=2)
        doc.add_paragraph(row.get("Story", ""))
        doc.add_paragraph(f"Source: {row.get('Source', '')}")
        doc.add_paragraph(f"Published At: {row.get('Published At', '')}")
        if row.get("URL"):
            doc.add_paragraph(f"URL: {row.get('URL')}")

    doc.save(str(path))


def _write_markdown(path: Path, rows: List[Dict[str, Any]], topic: str) -> None:
    lines = [f"# Top latest news for: {topic}", ""]
    for row in rows:
        lines.append(f"## {row.get('Rank', '')}. {row.get('Headline', '')}")
        lines.append("")
        lines.append(row.get("Story", ""))
        lines.append("")
        lines.append(f"- Source: {row.get('Source', '')}")
        lines.append(f"- Published At: {row.get('Published At', '')}")
        if row.get("URL"):
            lines.append(f"- URL: {row.get('URL')}")
        lines.append("")
    path.write_text("\n".join(lines).strip() + "\n", encoding="utf-8")


def _write_json(path: Path, rows: List[Dict[str, Any]]) -> None:
    path.write_text(json.dumps(rows, indent=2, ensure_ascii=False), encoding="utf-8")


def _write_csv(path: Path, rows: List[Dict[str, Any]], columns: List[str]) -> None:
    output = StringIO()
    writer = csv.DictWriter(output, fieldnames=columns)
    writer.writeheader()
    for row in rows:
        writer.writerow({col: row.get(col, "") for col in columns})
    path.write_text(output.getvalue(), encoding="utf-8", newline="")


@tool_registry.register(
    "create_latest_news_file",
    "Fetch latest news and write to file based on extension (.xlsx, .txt, .docx, .md, .json, .csv).",
    CreateLatestNewsFileArgs,
)
def create_latest_news_file(args: CreateLatestNewsFileArgs) -> str:
    try:
        data = _fetch_news_rows(
            args.topic,
            args.limit,
            args.feed_urls,
            args.sort_mode,
            args.story_mode,
        )
        rows = data["rows"]

        from app.sandbox.limits import check_file_size, resolve_sandbox_path

        path = resolve_sandbox_path(args.path)
        path.parent.mkdir(parents=True, exist_ok=True)
        ext = path.suffix.lower()

        columns = args.columns or ["Rank", "Headline", "Story", "Source", "Published At"]

        if ext == ".xlsx":
            _write_xlsx(path, rows, columns, args.sheet_name)
        elif ext == ".txt":
            _write_text(path, rows, args.topic)
        elif ext == ".docx":
            _write_docx(path, rows, args.topic)
        elif ext == ".md":
            _write_markdown(path, rows, args.topic)
        elif ext == ".json":
            _write_json(path, rows)
        elif ext == ".csv":
            _write_csv(path, rows, columns)
        else:
            return _err("Unsupported file format. Use .xlsx, .txt, .docx, .md, .json, or .csv")

        check_file_size(path)

        return _ok(
            f"Created {args.path} with {data['returned_count']} news rows.",
            data={
                "file_path": args.path,
                "format": ext,
                "requested_topic": args.topic,
                "normalized_topic": data["normalized_topic"],
                "requested_limit": args.limit,
                "returned_count": data["returned_count"],
                "shortage": data["shortage"],
                "rows": rows,
                "columns": columns,
                "partial_source_errors": data["partial_source_errors"],
                "story_enrichment_errors": data.get("story_enrichment_errors", 0),
                "story_source_counts": data.get("story_source_counts", {}),
                "story_depth": data.get("story_depth", {}),
            },
            files_modified=[args.path],
        )
    except Exception as exc:
        return _err(str(exc))
