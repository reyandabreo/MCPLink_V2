import json
import os
import re
from datetime import datetime
from pathlib import Path
from typing import Optional, List, Dict, Any

from docx import Document
from docx.shared import Pt, RGBColor
from openai import OpenAI
from pydantic import BaseModel, Field

from app.config import settings
from app.sandbox.limits import check_file_size, resolve_sandbox_path
from app.tools._llm_helper import call_llm
from app.tools.registry import tool_registry


def _ok(message: str, data: dict = None, files_modified: list = None) -> str:
    """Return success JSON response."""
    return json.dumps(
        {
            "status": "success",
            "message": message,
            "data": data or {},
            "files_modified": files_modified or [],
        }
    )


def _err(message: str) -> str:
    """Return error JSON response."""
    return json.dumps({"status": "error", "message": message})


def _parse_create_command(query: str) -> tuple[Optional[str], Optional[str]]:
    """
    Parse natural language command to extract filename and topic.
    Pattern: "create filename.docx with topic information"
    """
    pattern = r"\bcreate\s+(?:a\s+file\s+)?[\"']?([A-Za-z0-9_\-./ ]+\.(?:docx|txt|md|json|xlsx))[\"']?\s+(?:with|about|regarding|on)\s+(.+)"
    match = re.search(pattern, query or "", flags=re.IGNORECASE)
    
    if not match:
        return None, None
    
    file_path = match.group(1).strip()
    topic = match.group(2).strip().rstrip(".")
    return file_path, topic


def _generate_content_with_llm(query: str) -> str:
    """Get LLM response for the query."""
    try:
        response = call_llm(query, plain_text=True)
        return response.strip() if response else ""
    except Exception as e:
        return f"Error: {str(e)}"


def _generate_content_with_openai(topic: str) -> str:
    """
    Alternative: Generate content using OpenAI API directly.
    """
    try:
        api_key = os.getenv("OPENAI_API_KEY", settings.OPENAI_API_KEY if hasattr(settings, "OPENAI_API_KEY") else "")
        if not api_key or api_key == "YOUR_API_KEY_HERE":
            return f"[MOCK MODE]\n\nThis is a placeholder document regarding:\n{topic}\n\nConfigure OPENAI_API_KEY to generate real content."
        
        client = OpenAI(api_key=api_key)
        response = client.chat.completions.create(
            model="gpt-3.5-turbo",
            messages=[
                {"role": "system", "content": "You are a helpful assistant writing content for a document."},
                {"role": "user", "content": f"Write a detailed, well-structured summary about: {topic}"}
            ]
        )
        return response.choices[0].message.content.strip()
    except Exception as e:
        return f"Error generating content: {str(e)}"


def _parse_markdown_to_sections(content: str) -> List[Dict[str, Any]]:
    """
    Parse markdown content into structured sections for multi-format export.
    Returns list of section dicts with type, level, and content.
    """
    sections = []
    lines = content.split('\n')
    current_section = None
    current_content = []
    
    for line in lines:
        stripped = line.strip()
        
        if stripped.startswith('### '):
            if current_section:
                current_section['content'] = '\n'.join(current_content)
                sections.append(current_section)
            current_section = {'type': 'heading', 'level': 3, 'text': stripped[4:], 'content': ''}
            current_content = []
        elif stripped.startswith('## '):
            if current_section:
                current_section['content'] = '\n'.join(current_content)
                sections.append(current_section)
            current_section = {'type': 'heading', 'level': 2, 'text': stripped[3:], 'content': ''}
            current_content = []
        elif stripped.startswith('# '):
            if current_section:
                current_section['content'] = '\n'.join(current_content)
                sections.append(current_section)
            current_section = {'type': 'heading', 'level': 1, 'text': stripped[2:], 'content': ''}
            current_content = []
        elif stripped.startswith('- ') or stripped.startswith('* ') or stripped.startswith('• '):
            bullet_text = re.sub(r'^[-*•]\s*', '', stripped).strip()
            current_content.append(f'• {bullet_text}')
        elif re.match(r'^\d+\.\s+', stripped):
            numbered_text = re.sub(r'^\d+\.\s*', '', stripped).strip()
            current_content.append(f'◦ {numbered_text}')
        elif stripped:
            current_content.append(stripped)
        else:
            if current_content:
                current_content.append('')
    
    if current_section:
        current_section['content'] = '\n'.join(current_content)
        sections.append(current_section)
    elif current_content and any(line.strip() for line in current_content):
        # Ensure plain-text responses still produce at least one structured section.
        sections.append(
            {
                'type': 'heading',
                'level': 1,
                'text': 'Content',
                'content': '\n'.join(current_content).strip(),
            }
        )
    
    return sections


def _strip_markdown_formatting(text: str) -> str:
    """Remove markdown formatting markers for plain text output."""
    text = re.sub(r'\*\*\*(.+?)\*\*\*', r'\1', text)
    text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
    text = re.sub(r'__(.+?)__', r'\1', text)
    text = re.sub(r'\*(.+?)\*', r'\1', text)
    text = re.sub(r'_(.+?)_', r'\1', text)
    text = re.sub(r'`(.+?)`', r'\1', text)
    text = re.sub(r'~~(.+?)~~', r'\1', text)
    text = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'\1 (\2)', text)
    return text


def _write_txt_file(path: Path, title: str, content: str) -> None:
    """Write content to a plain text file with clean formatting."""
    lines = []
    lines.append(title)
    lines.append("=" * len(title))
    lines.append("")
    
    for line in content.split('\n'):
        stripped = line.strip()
        if stripped:
            clean_line = _strip_markdown_formatting(stripped)
            lines.append(clean_line)
        else:
            lines.append("")
    
    lines.append("")
    lines.append("-" * 50)
    lines.append(f"Generated on {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    
    path.write_text('\n'.join(lines), encoding='utf-8')


def _write_md_file(path: Path, title: str, content: str) -> None:
    """Write content to a markdown file (preserves markdown formatting)."""
    lines = []
    lines.append(f"# {title}")
    lines.append("")
    lines.append(content)
    lines.append("")
    lines.append("---")
    lines.append("")
    lines.append(f"*Generated on {datetime.now().strftime('%Y-%m-%d %H:%M')}*")
    
    path.write_text('\n'.join(lines), encoding='utf-8')


def _write_json_file(path: Path, title: str, content: str, topic: str) -> None:
    """Write content to a JSON file with structured metadata."""
    sections = _parse_markdown_to_sections(content)
    
    data = {
        "metadata": {
            "title": title,
            "topic": topic,
            "generated_at": datetime.now().isoformat(),
            "format_version": "1.0"
        },
        "content": {
            "raw": content,
            "sections": sections
        }
    }
    
    path.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding='utf-8')


def _write_xlsx_file(path: Path, title: str, content: str) -> None:
    """Write content to an Excel file with structured sections."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Content"[:31]
    
    # Styles
    title_font = Font(bold=True, size=14, color="000000")
    heading_font = Font(bold=True, size=12, color="003366")
    normal_font = Font(size=11, color="000000")
    center_alignment = Alignment(horizontal='center')
    left_alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
    
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    alt_fill = PatternFill(start_color="D6DCE4", end_color="D6DCE4", fill_type="solid")
    
    row = 1
    
    # Title
    ws.merge_cells(f'A{row}:B{row}')
    cell = ws[f'A{row}']
    cell.value = title
    cell.font = title_font
    cell.alignment = center_alignment
    row += 1
    
    ws[f'A{row}'] = "Generated:"
    ws[f'B{row}'] = datetime.now().strftime('%Y-%m-%d %H:%M')
    row += 2
    
    # Headers
    ws[f'A{row}'] = "Section"
    ws[f'B{row}'] = "Content"
    for col in ['A', 'B']:
        cell = ws[f'{col}{row}']
        cell.font = heading_font
        cell.fill = header_fill
        cell.alignment = center_alignment
        cell.border = thin_border
    row += 1
    
    # Content sections
    sections = _parse_markdown_to_sections(content)
    if not sections:
        sections = [
            {
                'type': 'heading',
                'level': 1,
                'text': 'Response',
                'content': content.strip() or 'No content generated.',
            }
        ]
    section_num = 1
    alt_row = False
    
    for section in sections:
        if section['type'] == 'heading':
            level_prefix = "#" * section['level']
            section_text = f"{level_prefix} {section['text']}"
        else:
            section_text = section.get('content', '')
        
        ws[f'A{row}'] = f"{section_num}. {section['text'] if section['type'] == 'heading' else 'Content'}"
        ws[f'B{row}'] = section.get('content', '') if section['type'] == 'heading' else section_text
        
        for col in ['A', 'B']:
            cell = ws[f'{col}{row}']
            cell.font = normal_font
            cell.alignment = left_alignment
            cell.border = thin_border
            if alt_row:
                cell.fill = alt_fill
        
        row += 1
        section_num += 1
        alt_row = not alt_row
    
    # Column widths
    ws.column_dimensions['A'].width = 25
    ws.column_dimensions['B'].width = 80
    
    wb.save(str(path))


def _add_styled_paragraph(paragraph, text: str) -> None:
    """
    Parse markdown-style inline formatting and apply to docx paragraph.
    Supports: **bold**, *italic*, `code`, ~~strikethrough~~, [links](url)
    """
    patterns = [
        (r'(\*\*\*|___)(.+?)\1', {'bold': True, 'italic': True}),
        (r'(\*\*|__)(.+?)\1', {'bold': True}),
        (r'(?<!\w)([*_])(?!\s)(.+?)(?<!\s)\1(?!\w)', {'italic': True}),
        (r'`([^`]+)`', {'code': True}),
        (r'~~(.+?)~~', {'strike': True}),
        (r'\[([^\]]+)\]\(([^)]+)\)', {'link': True}),
    ]
    
    def _process_segment(segment: str, styles: dict):
        """Apply formatting styles to a text segment."""
        run = paragraph.add_run(segment)
        
        if styles.get('bold'):
            run.font.bold = True
        if styles.get('italic'):
            run.font.italic = True
        if styles.get('strike'):
            run.font.strike = True
        if styles.get('code'):
            run.font.name = 'Consolas'
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0, 120, 0)
        if styles.get('link'):
            run.font.color.rgb = RGBColor(0, 102, 204)
            run.font.underline = True
        
        return run
    
    def _parse_inline(content: str, depth=0):
        """Recursively parse and add formatted runs."""
        if not content or depth > 3:
            if content.strip():
                run = paragraph.add_run(content)
                run.font.size = Pt(11)
                run.font.name = 'Calibri'
            return
        
        for pattern, styles in patterns:
            match = re.search(pattern, content)
            if match:
                start, end = match.span()
                
                if start > 0:
                    _parse_inline(content[:start], depth + 1)
                
                full_match = match.group(0)
                if styles.get('link'):
                    link_text = match.group(1)
                    _process_segment(link_text, {'bold': False, 'italic': False, 'link': True})
                elif styles.get('code'):
                    code_content = match.group(1)
                    _process_segment(code_content, {'code': True})
                else:
                    inner_content = match.group(2) if len(match.groups()) >= 2 else match.group(1)
                    _process_segment(inner_content, styles)
                
                if end < len(content):
                    _parse_inline(content[end:], depth + 1)
                return
        
        if content.strip():
            run = paragraph.add_run(content)
            run.font.size = Pt(11)
            run.font.name = 'Calibri'
    
    _parse_inline(text)


def _write_docx_file(path: Path, title: str, content: str) -> None:
    """Create a .docx file with properly formatted content."""
    doc = Document()
    
    # Add title as heading
    if title:
        heading = doc.add_heading(title, level=1)
        heading.runs[0].font.size = Pt(16)
        heading.runs[0].font.bold = True
    
    # Add separator
    separator = doc.add_paragraph()
    sep_run = separator.add_run("_" * 50)
    sep_run.font.size = Pt(8)
    
    # Process content line by line
    for line in content.split('\n'):
        stripped = line.strip()
        if stripped:
            if stripped.startswith('### '):
                doc.add_heading(stripped[4:], level=3)
            elif stripped.startswith('## '):
                doc.add_heading(stripped[3:], level=2)
            elif stripped.startswith('# '):
                doc.add_heading(stripped[2:], level=1)
            elif stripped.startswith('- ') or stripped.startswith('* ') or stripped.startswith('• '):
                para = doc.add_paragraph(style='List Bullet')
                bullet_text = re.sub(r'^[-*•]\s*', '', stripped).strip()
                _add_styled_paragraph(para, bullet_text)
            elif re.match(r'^\d+\.\s+', stripped):
                para = doc.add_paragraph(style='List Number')
                numbered_text = re.sub(r'^\d+\.\s*', '', stripped).strip()
                _add_styled_paragraph(para, numbered_text)
            else:
                para = doc.add_paragraph()
                _add_styled_paragraph(para, stripped)
    
    # Add footer
    doc.add_paragraph().add_run("_" * 50).font.size = Pt(8)
    footer = doc.add_paragraph()
    footer.alignment = 1  # Center
    footer_run = footer.add_run(f"Generated on {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    footer_run.font.size = Pt(9)
    footer_run.font.italic = True
    footer_run.font.color.rgb = RGBColor(100, 100, 100)
    
    doc.save(str(path))


def _write_file_by_extension(path: Path, title: str, content: str, topic: str) -> None:
    """Route content to appropriate writer based on file extension."""
    ext = path.suffix.lower()
    
    if ext == '.docx':
        _write_docx_file(path, title, content)
    elif ext == '.txt':
        _write_txt_file(path, title, content)
    elif ext == '.md':
        _write_md_file(path, title, content)
    elif ext == '.json':
        _write_json_file(path, title, content, topic)
    elif ext == '.xlsx':
        _write_xlsx_file(path, title, content)
    else:
        raise ValueError(f"Unsupported file extension: {ext}")


class WebFetchArgs(BaseModel):
    """Arguments for the web_fetch_tool."""
    query: str = Field(
        ..., 
        description="Natural language command like: 'create report.docx with information about ML libraries'"
    )
    output_path: Optional[str] = Field(
        None, 
        description="Optional explicit output path (overrides path in query)"
    )


@tool_registry.register(
    "web_fetch_tool",
    "Create a document (.docx, .txt, .md, .json, .xlsx) with AI-generated content based on a natural language query.",
    WebFetchArgs,
)
def web_fetch_tool(args: WebFetchArgs) -> str:
    """
    Get LLM response for user query and optionally save to file.
    
    Can be used in two ways:
    1. Just get an answer: query="Explain Python"
    2. Create a file: query="create python.docx with ML info" or query="Save to file" + output_path="file.txt"
    
    Supported formats: .docx, .txt, .md, .json, .xlsx
    """
    try:
        query = args.query.strip()
        if not query:
            return _err("Query cannot be empty")

        response = _generate_content_with_llm(query)
        if not response or response.startswith("Error"):
            return _err(f"Failed to generate content: {response}")

        parsed_path, parsed_topic = _parse_create_command(query)
        requested_path = (args.output_path or "").strip() or parsed_path

        if not requested_path:
            return _ok(
                "Response generated successfully",
                data={
                    "file_created": False,
                    "answer": response,
                },
            )

        ext = Path(requested_path).suffix.lower()
        if ext not in [".docx", ".txt", ".md", ".json", ".xlsx"]:
            return _err("Unsupported file format. Use .docx, .txt, .md, .json, or .xlsx")

        title = parsed_topic if parsed_topic else query
        if len(title) > 60:
            title = title[:60].rsplit(" ", 1)[0]

        out_path = resolve_sandbox_path(requested_path)
        out_path.parent.mkdir(parents=True, exist_ok=True)

        _write_file_by_extension(out_path, title, response, title)
        check_file_size(out_path)

        return _ok(
            f"Successfully created '{out_path.name}'",
            data={
                "file_created": True,
                "file_path": requested_path,
                "file_type": ext,
                "answer": response,
                "content_length": len(response),
            },
            files_modified=[requested_path],
        )

    except PermissionError as e:
        return _err(f"Permission denied: {str(e)}")
    except Exception as e:
        return _err(f"Unexpected error: {str(e)}")